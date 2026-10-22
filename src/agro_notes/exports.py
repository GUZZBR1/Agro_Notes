"""Read-only, property-scoped exports of the farmer-facing livestock records."""

from __future__ import annotations

import json
import math
import re
import unicodedata
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import Any, Protocol
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportStore(Protocol):
    def read_rows(self, tab: str) -> list[dict[str, Any]]: ...


@dataclass(frozen=True)
class ExportIntent:
    kind: str = "complete"


@dataclass(frozen=True)
class ExportDocument:
    filename: str
    content: bytes
    mime_type: str = XLSX_MIME_TYPE


class _MemoryExcelWriter(ExcelWriter):
    """Use openpyxl's supported XML stream without its default disk scratch files.

    ExcelWriter writes worksheet paths by default. Keep this adaptation confined to
    this writer instance so concurrent exports never mutate process-wide state.
    """

    def write_worksheet(self, worksheet: Any) -> None:
        worksheet._drawing = SpreadsheetDrawing()
        worksheet._drawing.charts = worksheet._charts
        worksheet._drawing.images = worksheet._images
        with BytesIO() as stream:
            writer = WorksheetWriter(worksheet, out=stream)
            try:
                writer.write()
                worksheet._rels = writer._rels
                self._archive.writestr(worksheet.path[1:], stream.getvalue())
                self.manifest.append(worksheet)
            finally:
                writer.close()


def _save_workbook(workbook: Workbook) -> bytes:
    with BytesIO() as output:
        with ZipFile(output, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            _MemoryExcelWriter(workbook, archive).save()
        return output.getvalue()


def match_export_intent(text: str) -> ExportIntent | None:
    """Recognize an explicit download request without interpreting property selectors."""
    normalized = "".join(
        char for char in unicodedata.normalize("NFKD", text.casefold()) if not unicodedata.combining(char)
    )
    words = set(re.findall(r"[a-z0-9]+", normalized))
    if re.match(r"^\s*/(?:exportar|export|planilha)(?:@[a-z0-9_]+)?(?:\s|$)", normalized):
        return ExportIntent()
    export = bool(
        words & {"exporta", "exportar", "exporte", "baixar", "baixa", "baixe", "download", "gera", "gere", "gerar"}
    )
    send = bool(words & {"manda", "mande", "mandar", "envia", "envie", "enviar"})
    subject = bool(words & {"planilha", "excel", "xlsx", "dados", "historico", "registros"})
    return ExportIntent() if subject and (export or send) else None


def _scalar(value: Any) -> str | int | float | bool:
    # Containers are never serialized: they could contain hidden IDs or raw metadata.
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
    if isinstance(value, bool | int):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    return ""


def _latest(store: ExportStore, tab: str, property_id: str, key: str) -> list[dict[str, Any]]:
    scoped: dict[str, dict[str, Any]] = {}
    for row in store.read_rows(tab):
        # Filter BEFORE deduplication or joining; blank property IDs are not trusted.
        if row.get("property_id") != property_id:
            continue
        identity = row.get(key)
        if not isinstance(identity, str) or not identity:
            continue
        scoped[identity] = row
    return list(scoped.values())


def _payload(event: Mapping[str, Any]) -> dict[str, Any]:
    try:
        value = json.loads(event.get("event_payload_json") or "{}")
    except (TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def _belongs(value: Any, entities: Mapping[str, Any]) -> bool:
    return isinstance(value, str) and value in entities


EVENT_LABELS = {
    "ANIMAL_CREATED": "Cadastro de animal",
    "ANIMAL_ENTERED": "Entrada de animal",
    "CALVING": "Parto",
    "WEIGHING": "Pesagem",
    "LOT_TRANSFER": "Transferência de lote",
    "VACCINATION": "Vacinação reportada",
    "HEALTH_OBSERVATION": "Observação de saúde",
}

# Explicit columns are the privacy boundary. No raw rows or payload JSON reach XLSX.
ANIMAL_FIELDS = (
    ("species", "Espécie"),
    ("sex", "Sexo"),
    ("birth_date", "Nascimento"),
    ("birth_date_precision", "Precisão do nascimento"),
    ("breed", "Raça"),
    ("origin_type", "Origem"),
    ("first_entry_date", "Entrada na fazenda"),
)
EVENT_FIELDS = {
    "ANIMAL_CREATED": (("species", "Espécie"), ("sex", "Sexo"), ("breed", "Raça")),
    "ANIMAL_ENTERED": (("origin", "Origem"),),
    "WEIGHING": (("weight_kg", "Peso (kg)"),),
    "VACCINATION": (("vaccine", "Vacina"), ("dose", "Dose")),
    "HEALTH_OBSERVATION": (
        ("observation", "Observação"),
        ("observed_sign", "Sinal observado"),
        ("body_region", "Região do corpo"),
    ),
    "CALVING": (("outcome_count", "Quantidade de crias"),),
    "LOT_TRANSFER": (),
}


def build_property_export(store: ExportStore, property_id: str, kind: str = "complete") -> ExportDocument:
    """Build XLSX entirely in memory; caller supplies its already-authorized property."""
    if not isinstance(property_id, str) or not property_id.strip():
        raise ValueError("A trusted property_id is required")
    if kind != "complete":
        raise ValueError("Unsupported export kind")
    animals = {row["animal_id"]: row for row in _latest(store, "02_Animais", property_id, "animal_id")}
    groups = {row["group_id"]: row for row in _latest(store, "04_Grupos", property_id, "group_id")}
    identifiers = _latest(store, "03_Identificadores_Animal", property_id, "identifier_id")
    allocations = _latest(store, "05_Alocacoes", property_id, "allocation_id")
    events = {row["event_id"]: row for row in _latest(store, "06_Eventos", property_id, "event_id")}
    targets = _latest(store, "07_Evento_Animais", property_id, "event_target_id")

    labels: dict[str, str] = {}
    for row in identifiers:
        animal_id = row.get("animal_id")
        if (
            _belongs(animal_id, animals)
            and not row.get("valid_to")
            and str(row.get("status", "")).upper() not in {"INACTIVE", "REVOKED", "REPLACED"}
            and isinstance(row.get("value"), str | int | float)
        ):
            labels[animal_id] = str(row["value"])

    def animal_label(animal_id: Any) -> str:
        return labels.get(animal_id, "Sem identificador") if isinstance(animal_id, str) and animal_id in animals else ""

    def group_label(group_id: Any) -> str:
        if not isinstance(group_id, str) or group_id not in groups:
            return ""
        name = groups[group_id].get("name")
        return name if isinstance(name, str) else ""

    current_groups: dict[str, list[str]] = {}
    for row in allocations:
        if (
            _belongs(row.get("animal_id"), animals)
            and _belongs(row.get("group_id"), groups)
            and not row.get("effective_to")
        ):
            current_groups.setdefault(row["animal_id"], []).append(str(group_label(row["group_id"])))

    event_animals: dict[str, set[str]] = {}
    event_groups: dict[str, set[str]] = {}
    for row in targets:
        event_id = row.get("event_id")
        if not _belongs(event_id, events):
            continue
        if _belongs(row.get("animal_id"), animals):
            event_animals.setdefault(event_id, set()).add(animal_label(row["animal_id"]))
        if _belongs(row.get("group_id"), groups):
            event_groups.setdefault(event_id, set()).add(str(group_label(row["group_id"])))

    # A successful correction replaces its predecessor in current reports.
    corrected = {
        row["correction_of"]
        for row in events.values()
        if _belongs(row.get("correction_of"), events) and row.get("processing_status") == "APPLIED"
    }
    active_events = [
        row
        for row in events.values()
        if row["event_id"] not in corrected
        and row.get("event_type") in EVENT_LABELS
        and row.get("processing_status") == "APPLIED"
    ]
    active_events.sort(key=lambda row: (str(row.get("occurred_from", "")), row["event_id"]))

    sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    if animals:
        animal_rows = [
            [
                animal_label(animal_id),
                *[row.get(field, "") for field, _ in ANIMAL_FIELDS],
                "; ".join(sorted(set(current_groups.get(animal_id, [])))),
            ]
            for animal_id, row in sorted(animals.items(), key=lambda pair: (animal_label(pair[0]), pair[0]))
        ]
        sheets.append(("Animais", ["Animal", *[label for _, label in ANIMAL_FIELDS], "Lote atual"], animal_rows))
    histories: list[list[Any]] = []
    typed_rows: dict[str, list[list[Any]]] = {key: [] for key in ("WEIGHING", "VACCINATION", "HEALTH_OBSERVATION")}
    for event in active_events:
        event_type = event["event_type"]
        payload = _payload(event)
        event_id = event["event_id"]
        animal_names = set(event_animals.get(event_id, set()))
        group_names = set(event_groups.get(event_id, set()))
        # Payload references are resolved exclusively against the scoped entities.
        name = animal_label(payload.get("animal_id") or payload.get("mother_animal_id"))
        if name:
            animal_names.add(name)
        name = group_label(payload.get("group_id"))
        if name:
            group_names.add(str(name))
        common = [event.get("occurred_from", ""), "; ".join(sorted(animal_names)), "; ".join(sorted(group_names))]
        detail = "; ".join(
            f"{label}: {_scalar(payload[field])}"
            for field, label in EVENT_FIELDS[event_type]
            if _scalar(payload.get(field)) != ""
        )
        if event_type == "LOT_TRANSFER":
            detail = "; ".join(
                f"{label}: {group_label(payload.get(field))}"
                for field, label in (("from_group_id", "Origem"), ("to_group_id", "Destino"))
                if group_label(payload.get(field))
            )
        histories.append([common[0], EVENT_LABELS[event_type], common[1], common[2], detail])
        if event_type in typed_rows:
            typed_rows[event_type].append([*common, *[payload.get(field, "") for field, _ in EVENT_FIELDS[event_type]]])
    for event_type, title in (("WEIGHING", "Pesagens"), ("VACCINATION", "Vacinação"), ("HEALTH_OBSERVATION", "Saúde")):
        if typed_rows[event_type]:
            sheets.append(
                (
                    title,
                    ["Data", "Animal", "Lote", *[label for _, label in EVENT_FIELDS[event_type]]],
                    typed_rows[event_type],
                )
            )
    if groups:
        lot_rows = [
            [
                row.get("name", ""),
                row.get("group_type", ""),
                row.get("purpose", ""),
                row.get("valid_from", ""),
                row.get("valid_to", ""),
                sum(
                    group_id == allocation.get("group_id")
                    and _belongs(allocation.get("animal_id"), animals)
                    and not allocation.get("effective_to")
                    for allocation in allocations
                ),
            ]
            for group_id, row in sorted(groups.items(), key=lambda pair: (str(pair[1].get("name", "")), pair[0]))
        ]
        sheets.append(("Lotes", ["Lote", "Tipo", "Finalidade", "Início", "Fim", "Alocações abertas"], lot_rows))
    if histories:
        sheets.append(("Histórico", ["Data", "Evento", "Animal", "Lote", "Detalhes"], histories))

    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = [["Animais", len(animals)], ["Lotes", len(groups)], ["Eventos", len(active_events)]]
    if not animals and not groups and not active_events:
        summary.append(["Situação", "Ainda não há dados registrados nesta fazenda."])
    sheets.insert(0, ("Resumo", ["Informação", "Valor"], summary))
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(title)
        for values in [headers, *rows]:
            sheet.append([_scalar(value) for value in values])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    # XLSX explicit string cells prevent formulas, including =, +, -, @.
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245A3D")
        for column in range(1, len(headers) + 1):
            width = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1))
            sheet.column_dimensions[get_column_letter(column)].width = min(55, max(14, width + 2))
    safe_property = re.sub(r"[^A-Za-z0-9_-]", "_", property_id)[:64].strip("_-") or "fazenda"
    filename = f"agro_notes_{safe_property}_{datetime.now(UTC).date().isoformat()}.xlsx"
    return ExportDocument(filename=filename, content=_save_workbook(workbook))
