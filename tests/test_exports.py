import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from agro_notes.exports import XLSX_MIME_TYPE, build_property_export, match_export_intent


class ExportStore:
    def __init__(self):
        self.tabs = {}
        self.reads = []

    def read_rows(self, tab):
        self.reads.append(tab)
        assert tab not in {"18_Conversas", "19_Usuarios", "20_Event_Effects", "99_Audit_Log", "Telegram_Access"}
        return self.tabs.get(tab, [])


def workbook(store):
    document = build_property_export(store, "farm")
    assert document.filename.endswith(".xlsx")
    assert document.mime_type == XLSX_MIME_TYPE
    assert document.content.startswith(b"PK")
    return load_workbook(BytesIO(document.content))


def records(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]


@pytest.mark.parametrize(
    "text",
    [
        "me manda minha planilha",
        "exporta meus dados",
        "me manda o histórico da fazenda em Excel",
        "Pode enviar uma planilha?",
        "quero baixar meus registros",
        "/exportar",
        "/exportar@AgroBot",
        "gera um Excel",
    ],
)
def test_matches_explicit_export_requests(text):
    assert match_export_intent(text) is not None


@pytest.mark.parametrize(
    "text",
    [
        "a vaca 42 pesou 300 kg",
        "planilha",
        "O que significa exportação?",
        "mandarina",
        "Excel é útil",
        "me manda a vacina",
        "/exportar_animais",
        "o animal chama Excel",
    ],
)
def test_does_not_route_ordinary_capture_or_explanation_as_export(text):
    assert match_export_intent(text) is None


def test_empty_workbook_is_valid_and_does_not_read_sensitive_tabs():
    store = ExportStore()
    result = workbook(store)
    assert result.sheetnames == ["Resumo"]
    assert records(result["Resumo"])[-1]["Valor"] == "Ainda não há dados registrados nesta fazenda."


@pytest.mark.parametrize("property_id", ["", " ", None])
def test_missing_trusted_property_fails_before_reading(property_id):
    store = ExportStore()
    with pytest.raises(ValueError, match="trusted property_id"):
        build_property_export(store, property_id)
    assert store.reads == []


def test_filtered_before_deduplication_and_joins_and_no_raw_metadata():
    store = ExportStore()
    store.tabs = {
        "02_Animais": [
            {"property_id": "farm", "animal_id": "a", "species": "Bovino", "secret": "TOKEN_SENTINEL"},
            {"property_id": "other", "animal_id": "a", "species": "FOREIGN_SENTINEL"},
            {"property_id": "other", "animal_id": "b", "species": "OTHER_ANIMAL"},
        ],
        "03_Identificadores_Animal": [
            {"property_id": "farm", "identifier_id": "i", "animal_id": "a", "value": "42"},
            {"property_id": "other", "identifier_id": "i", "animal_id": "a", "value": "SECRET_LABEL"},
            {"property_id": "", "identifier_id": "legacy", "animal_id": "a", "value": "UNSCOPED_LABEL"},
            {"property_id": "farm", "identifier_id": "orphan", "animal_id": "b", "value": "ORPHAN_LABEL"},
        ],
        "04_Grupos": [
            {"property_id": "farm", "group_id": "g", "name": "Lote 1"},
            {"property_id": "other", "group_id": "foreign", "name": "FOREIGN_GROUP"},
        ],
        "06_Eventos": [
            {
                "property_id": "farm",
                "event_id": "e",
                "event_type": "WEIGHING",
                "processing_status": "APPLIED",
                "occurred_from": "2026-10-12",
                "actor_id": "PRIVATE_ACTOR",
                "command_id": "PRIVATE_COMMAND",
                "event_payload_json": json.dumps({"weight_kg": 300, "animal_id": "a", "token": "SECRET_TOKEN"}),
            },
            {
                "property_id": "other",
                "event_id": "e",
                "event_type": "VACCINATION",
                "event_payload_json": '{"vaccine":"PRIVATE_VACCINE"}',
            },
        ],
        "07_Evento_Animais": [
            {"property_id": "farm", "event_target_id": "t", "event_id": "e", "animal_id": "a"},
            {"property_id": "farm", "event_target_id": "bad", "event_id": "e", "animal_id": "b", "group_id": "foreign"},
            {"property_id": "other", "event_target_id": "t", "event_id": "e", "animal_id": "b"},
        ],
    }
    result = workbook(store)
    assert records(result["Animais"])[0]["Animal"] == "42"
    assert records(result["Animais"])[0]["Espécie"] == "Bovino"
    assert records(result["Pesagens"])[0]["Peso (kg)"] == 300
    assert records(result["Pesagens"])[0]["Animal"] == "42"
    assert records(result["Pesagens"])[0]["Lote"] is None
    serialized = str([list(sheet.values) for sheet in result])
    for secret in ("SENTINEL", "SECRET_", "PRIVATE_", "FOREIGN_", "ORPHAN_", "UNSCOPED_", "animal_id", "command_id"):
        assert secret not in serialized


def test_latest_versions_close_allocations_revoke_identifiers_and_replace_corrected_events():
    store = ExportStore()
    store.tabs = {
        "02_Animais": [{"property_id": "farm", "animal_id": "a"}],
        "04_Grupos": [{"property_id": "farm", "group_id": "g", "name": "Lote antigo"}],
        "03_Identificadores_Animal": [
            {"property_id": "farm", "identifier_id": "i", "animal_id": "a", "value": "OLD", "status": "ACTIVE"},
            {"property_id": "farm", "identifier_id": "i", "animal_id": "a", "value": "OLD", "status": "REVOKED"},
            {"property_id": "farm", "identifier_id": "j", "animal_id": "a", "value": "43", "status": "ACTIVE"},
        ],
        "05_Alocacoes": [
            {"property_id": "farm", "allocation_id": "allocation", "animal_id": "a", "group_id": "g"},
            {
                "property_id": "farm",
                "allocation_id": "allocation",
                "animal_id": "a",
                "group_id": "g",
                "effective_to": "2026-10-12",
            },
        ],
        "06_Eventos": [
            {
                "property_id": "farm",
                "event_id": "e",
                "event_type": "WEIGHING",
                "processing_status": "PENDING",
                "event_payload_json": '{"weight_kg":300}',
            },
            {
                "property_id": "farm",
                "event_id": "e",
                "event_type": "WEIGHING",
                "processing_status": "APPLIED",
                "event_payload_json": '{"weight_kg":301}',
            },
            {
                "property_id": "farm",
                "event_id": "fixed",
                "event_type": "WEIGHING",
                "processing_status": "APPLIED",
                "correction_of": "e",
                "event_payload_json": '{"weight_kg":302,"animal_id":"a"}',
            },
        ],
    }
    result = workbook(store)
    assert records(result["Animais"])[0]["Animal"] == "43"
    assert records(result["Animais"])[0]["Lote atual"] is None
    assert records(result["Lotes"])[0]["Alocações abertas"] == 0
    assert [row["Peso (kg)"] for row in records(result["Pesagens"])] == [302]
    assert records(result["Resumo"])[2]["Valor"] == 1


@pytest.mark.parametrize("value", ['=HYPERLINK("https://evil")', "+cmd", "-cmd", "@SUM(1)", "\t=1+1", "\r=1+1"])
def test_untrusted_values_are_literal_cells_never_formulas(value):
    store = ExportStore()
    store.tabs = {
        "02_Animais": [{"property_id": "farm", "animal_id": "a", "species": value}],
        "03_Identificadores_Animal": [{"property_id": "farm", "identifier_id": "i", "animal_id": "a", "value": value}],
        "04_Grupos": [{"property_id": "farm", "group_id": "g", "name": value}],
        "06_Eventos": [
            {
                "property_id": "farm",
                "event_id": "e",
                "event_type": "HEALTH_OBSERVATION",
                "processing_status": "APPLIED",
                "event_payload_json": json.dumps({"observation": value}),
            }
        ],
    }
    result = workbook(store)
    for sheet in result:
        for row in sheet:
            assert all(cell.data_type != "f" and cell.hyperlink is None for cell in row)
    assert records(result["Saúde"])[0]["Observação"].replace("\n", "\r") == value.replace("\n", "\r")


@pytest.mark.parametrize(
    "payload", ["not json", "[]", '"text"', "null", '{"animal_id":[],"observation":{"token":"hidden"}}']
)
def test_malformed_payload_is_not_exposed_and_cannot_break_export(payload):
    store = ExportStore()
    store.tabs["06_Eventos"] = [
        {
            "property_id": "farm",
            "event_id": "e",
            "event_type": "HEALTH_OBSERVATION",
            "processing_status": "APPLIED",
            "event_payload_json": payload,
        }
    ]
    result = workbook(store)
    assert "hidden" not in str(list(result["Histórico"].values))
    assert records(result["Saúde"])[0]["Observação"] is None


def test_partial_events_are_excluded_without_internal_processing_fields():
    store = ExportStore()
    store.tabs["06_Eventos"] = [
        {
            "property_id": "farm",
            "event_id": "e",
            "event_type": "VACCINATION",
            "processing_status": "PARTIALLY_APPLIED",
            "event_payload_json": '{"vaccine":"Raiva","dose":"2 ml"}',
        }
    ]
    result = workbook(store)
    assert result.sheetnames == ["Resumo"]
    assert records(result["Resumo"])[2]["Valor"] == 0


def test_in_memory_generation_never_allocates_worksheet_scratch_files(monkeypatch):
    def fail_if_scratch_file(*args, **kwargs):
        raise AssertionError("worksheet XML must remain in memory")

    monkeypatch.setattr("openpyxl.worksheet._writer.create_temporary_file", fail_if_scratch_file)
    assert workbook(ExportStore()).sheetnames == ["Resumo"]


def test_injected_serialization_failure_creates_no_scratch_files_and_can_retry(monkeypatch):
    from zipfile import ZipFile

    def fail_if_scratch_file(*args, **kwargs):
        raise AssertionError("worksheet XML must remain in memory")

    monkeypatch.setattr("openpyxl.worksheet._writer.create_temporary_file", fail_if_scratch_file)
    original_write = ZipFile.writestr

    def fail_worksheet(self, name, *args, **kwargs):
        if isinstance(name, str) and name.startswith("xl/worksheets/"):
            raise OSError("injected archive failure")
        return original_write(self, name, *args, **kwargs)

    monkeypatch.setattr(ZipFile, "writestr", fail_worksheet)
    with pytest.raises(OSError, match="injected archive failure"):
        build_property_export(ExportStore(), "farm")
    monkeypatch.setattr(ZipFile, "writestr", original_write)
    assert workbook(ExportStore()).sheetnames == ["Resumo"]


def test_filename_sanitizes_property_without_path_traversal():
    document = build_property_export(ExportStore(), "../../farm\\private")
    assert document.filename.startswith("agro_notes_farm_private_")
    assert "/" not in document.filename and "\\" not in document.filename


def test_container_group_names_are_never_stringified_into_export():
    store = ExportStore()
    store.tabs = {
        "02_Animais": [{"property_id": "farm", "animal_id": "a"}],
        "03_Identificadores_Animal": [{"property_id": "farm", "identifier_id": "i", "animal_id": "a", "value": "42"}],
        "04_Grupos": [
            {"property_id": "farm", "group_id": "g", "name": {"token": "SECRET_TOKEN", "group_id": "internal"}}
        ],
        "05_Alocacoes": [{"property_id": "farm", "allocation_id": "al", "animal_id": "a", "group_id": "g"}],
        "07_Evento_Animais": [
            {"property_id": "farm", "event_target_id": "t", "event_id": "e", "animal_id": "a", "group_id": "g"}
        ],
        "06_Eventos": [
            {"property_id": "farm", "event_id": "e", "event_type": "WEIGHING", "processing_status": "APPLIED",
             "event_payload_json": '{"weight_kg":300,"group_id":"g","animal_id":"a"}'}
        ],
    }
    result = workbook(store)
    serialized = str([list(sheet.values) for sheet in result])
    assert "SECRET_TOKEN" not in serialized
    assert "internal" not in serialized
    assert records(result["Animais"])[0]["Lote atual"] is None
    assert records(result["Pesagens"])[0]["Lote"] is None
