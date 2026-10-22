import asyncio
from copy import deepcopy
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from agro_notes.application import MVPProcessor
from agro_notes.domain import EventProposal, InterpretationProposal
from agro_notes.exports import build_property_export
from agro_notes.sheets import TAB_HEADERS


class MemoryStore:
    def __init__(self, *, fail_target=False):
        self.tabs = {name: [] for name in TAB_HEADERS}
        self.fail_target = fail_target
        self.ambiguous_identifier = None
        self.tabs["02_Animais"].append({"animal_id": "animal-42", "property_id": "farm"})
        self.tabs["03_Identificadores_Animal"].append(
            {
                "animal_id": "animal-42",
                "property_id": "farm",
                "value": "42",
                "valid_to": "",
                "status": "ACTIVE",
            }
        )

    def read_rows(self, tab):
        return list(self.tabs[tab])

    def append_rows(self, tab, rows):
        self.tabs[tab].extend(dict(row) for row in rows)

    def resolve_animal(self, property_id, identifier):
        if identifier == self.ambiguous_identifier:
            from agro_notes.sheets import AmbiguousResolutionError

            raise AmbiguousResolutionError(identifier)
        animals = {row["animal_id"]: row for row in self.tabs["02_Animais"] if row.get("property_id") == property_id}
        if identifier in animals:
            return animals[identifier]
        for row in self.tabs["03_Identificadores_Animal"]:
            if row.get("property_id") == property_id and row.get("value") == identifier:
                return animals.get(row["animal_id"])
        return None

    def resolve_group(self, property_id, identifier):
        return next(
            (
                row
                for row in self.tabs["04_Grupos"]
                if row.get("property_id") == property_id
                and identifier in (row.get("name"), row.get("group_id"))
                and not row.get("valid_to")
            ),
            None,
        )

    def find_event_by_command(self, command_id):
        return next((row for row in reversed(self.tabs["06_Eventos"]) if row.get("command_id") == command_id), None)

    def find_event(self, event_id):
        return next((row for row in reversed(self.tabs["06_Eventos"]) if row.get("event_id") == event_id), None)

    def find_effect(self, key):
        return next((row for row in reversed(self.tabs["20_Event_Effects"]) if row.get("idempotency_key") == key), None)

    def append_event(self, row):
        self.tabs["06_Eventos"].append(dict(row))

    def append_target(self, row):
        if self.fail_target:
            raise OSError("sheets unavailable")
        self.tabs["07_Evento_Animais"].append(dict(row))

    def append_effect(self, row):
        self.tabs["20_Event_Effects"].append(dict(row))

    def append_audit(self, row):
        self.tabs["99_Audit_Log"].append(dict(row))

    def append_conversation(self, row):
        self.tabs["18_Conversas"].append(dict(row))

    def append_calving_outcome(self, row):
        self.tabs["08A_Resultados_Parto"].append(dict(row))

    def append_animal(self, row):
        self.tabs["02_Animais"].append(dict(row))

    def append_identifier(self, row):
        self.tabs["03_Identificadores_Animal"].append(dict(row))

    def append_genealogy(self, row):
        self.tabs["02A_Vinculos_Genealogicos"].append(dict(row))

    def append_allocation(self, row):
        self.tabs["05_Alocacoes"].append(dict(row))


class FixedInterpreter:
    def __init__(self, proposal):
        self.proposal = proposal

    def interpret(self, _text):
        return self.proposal


class SequenceInterpreter:
    def __init__(self, *proposals):
        self.proposals = iter(proposals)

    def interpret(self, _text):
        return next(self.proposals)


def local_settings():
    return SimpleNamespace(
        allowed_user_ids=frozenset({42}),
        farm_property_id="farm",
        farm_timezone="America/Sao_Paulo",
        min_interpretation_confidence=0.75,
        interpretation_model="test",
        transcription_model="none",
    )


def make_processor(
    store, event_type="HEALTH_OBSERVATION", payload=None, correction_of=None, correction_requested=False
):
    payload = payload or {"animal_identifier": "42", "observation": "manca"}
    proposal = InterpretationProposal(
        events=[
            EventProposal(
                event_type=event_type,
                payload=payload,
                time_expression="hoje",
                confidence=0.98,
                correction_of=correction_of,
                correction_requested=correction_requested,
            )
        ]
    )
    settings = local_settings()
    return MVPProcessor(settings, store, FixedInterpreter(proposal))


def receive(processor, message_id="chat:1", source="text"):
    return processor.receive(
        42,
        "relato",
        source=source,
        message_id=message_id,
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )




def test_trusted_context_rejects_invalid_or_timezone_free_capture_time():
    processor = make_processor(MemoryStore())

    with pytest.raises(ValueError, match="captured_at inválido"):
        processor._trusted_context(42, "command", "not-a-date")
    with pytest.raises(ValueError, match="precisa incluir fuso"):
        processor._trusted_context(42, "command", "2026-10-13T10:00:00")


@pytest.mark.asyncio
async def test_ready_health_observation_appends_event_target_and_audit_once():
    store = MemoryStore()
    processor = make_processor(store)

    first = await receive(processor)
    repeated = await receive(processor)

    assert first.text.startswith("✅")
    assert "manca" in first.text and "sem diagnóstico" in first.text
    assert "já foi registrado" in repeated.text
    assert len(store.tabs["06_Eventos"]) == 2
    assert store.tabs["06_Eventos"][-1]["processing_status"] == "APPLIED"
    assert store.tabs["06_Eventos"][-1]["original_time_expression"] == "hoje"
    assert store.tabs["06_Eventos"][-1]["event_status"] == "RECORDED"
    assert store.tabs["06_Eventos"][-1]["interpretation_confirmation"] == "NOT_REQUIRED"
    assert store.tabs["07_Evento_Animais"][0]["animal_id"] == "animal-42"
    assert store.tabs["20_Event_Effects"][-1]["status"] == "APPLIED"
    assert any(row["action"] == "DUPLICATE_DELIVERY" for row in store.tabs["99_Audit_Log"])






@pytest.mark.asyncio
async def test_concurrent_duplicate_delivery_serializes_to_one_event():
    store = MemoryStore()
    processor = make_processor(store)

    results = await asyncio.gather(receive(processor), receive(processor))

    assert sum("Observação reportada" in result.text for result in results) == 1
    assert sum("já foi registrado" in result.text for result in results) == 1
    assert len(store.tabs["06_Eventos"]) == 2
    assert len(store.tabs["07_Evento_Animais"]) == 1






@pytest.mark.asyncio
async def test_confirmation_is_persisted_then_applies_and_marks_conversation():
    store = MemoryStore()
    processor = make_processor(store, "VACCINATION", {"animal_identifier": "42", "vaccine": "vacina X"})

    preview = await receive(processor)
    result = await processor.confirm(42, preview.pending_id)

    assert preview.pending_id == "LOCAL:42:chat:1"
    assert "Confira a prévia" in preview.text
    assert result.text.startswith("✅")
    assert store.tabs["06_Eventos"][-1]["interpretation_confirmation"] == "USER_CONFIRMED"
    assert store.tabs["06_Eventos"][-1]["event_status"] == "CONFIRMED"
    assert store.tabs["18_Conversas"][-1]["state"] == "COMPLETED"










@pytest.mark.asyncio
async def test_vaccination_without_target_asks_for_animal_or_group():
    store = MemoryStore()
    processor = make_processor(store, "VACCINATION", {"vaccine": "aftosa"})

    result = await receive(processor)

    assert "animal_identifier ou group_identifier" in result.text
    assert not store.tabs["06_Eventos"]


@pytest.mark.asyncio
async def test_group_vaccination_is_confirmed_and_appears_in_member_history():
    store = MemoryStore()
    store.tabs["04_Grupos"].append(
        {"group_id": "group-7", "name": "7", "property_id": "farm", "valid_to": ""}
    )
    store.tabs["05_Alocacoes"].append(
        {
            "allocation_id": "allocation-42",
            "property_id": "farm",
            "animal_id": "animal-42",
            "group_id": "group-7",
            "effective_from": "2026-01-01",
            "effective_to": "",
        }
    )
    store.tabs["02_Animais"].append({"animal_id": "animal-43", "property_id": "farm"})
    store.tabs["03_Identificadores_Animal"].append(
        {
            "animal_id": "animal-43",
            "property_id": "farm",
            "value": "43",
            "valid_to": "",
            "status": "ACTIVE",
        }
    )
    processor = make_processor(store, "VACCINATION", {"group_identifier": "7", "vaccine": "aftosa"})

    preview = await receive(processor)
    result = await processor.confirm(42, preview.pending_id)
    history = await processor.receive(
        42,
        "/historico 42",
        source="text",
        message_id="chat:group-history",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )
    other_history = await processor.receive(
        42,
        "/historico 43",
        source="text",
        message_id="chat:group-history-other",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )

    event = store.tabs["06_Eventos"][-1]
    target = store.tabs["07_Evento_Animais"][-1]
    assert "Lote: 7" in preview.text
    assert result.text.startswith("✅") and "lote 7" in result.text
    assert event["target_scope"] == "GROUP"
    assert target["group_id"] == "group-7" and target["animal_id"] == ""
    assert "Vacinação reportada no lote 7" in history.text
    assert "Vacinação" not in other_history.text


@pytest.mark.asyncio
async def test_partial_effect_failure_is_recorded_without_success_claim():
    store = MemoryStore(fail_target=True)
    processor = make_processor(store)

    result = await receive(processor)

    assert result.text.startswith("⚠️")
    assert store.tabs["20_Event_Effects"][-1]["status"] == "FAILED"
    assert store.tabs["06_Eventos"][-1]["processing_status"] == "PARTIALLY_APPLIED"
    assert store.tabs["18_Conversas"][-1]["state"] == "PARTIALLY_APPLIED"
    event_rows_before_retry = len(store.tabs["06_Eventos"])
    target_rows_before_retry = len(store.tabs["07_Evento_Animais"])

    retry = await receive(processor)

    assert "reconciliação" in retry.text
    assert len(store.tabs["06_Eventos"]) == event_rows_before_retry
    assert len(store.tabs["07_Evento_Animais"]) == target_rows_before_retry


@pytest.mark.asyncio
async def test_history_shows_partial_event_even_when_target_effect_failed():
    store = MemoryStore(fail_target=True)
    processor = make_processor(store)
    await receive(processor)

    history = await processor.receive(
        42,
        "/historico 42",
        source="text",
        message_id="chat:history-partial",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )

    assert "manca" in history.text
    assert "registro parcial" in history.text


@pytest.mark.asyncio
async def test_history_explains_ambiguous_animal_identifier():
    store = MemoryStore()
    store.ambiguous_identifier = "42"
    processor = make_processor(store)

    history = await processor.receive(
        42,
        "/historico 42",
        source="text",
        message_id="chat:history-ambiguous",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )

    assert "mais de um animal" in history.text
    assert "identificador único" in history.text


@pytest.mark.asyncio
async def test_animal_created_and_entered_are_distinct_canonical_events():
    created_store = MemoryStore()
    created = make_processor(
        created_store,
        "ANIMAL_CREATED",
        {
            "display_identifier": "77",
            "species": "BOVINE",
            "sex": "FEMALE",
            "breed": "Nelore",
            "origin": "compra",
            "birth_date": "2024-01-02",
        },
    )
    preview = await receive(created)
    assert "Espécie: BOVINE" in preview.text
    assert "Sexo: FEMALE" in preview.text
    assert "Raça: Nelore" in preview.text
    assert "Nascimento: 02/01/2024" in preview.text
    await created.confirm(42, preview.pending_id)
    created_animal = next(row for row in created_store.tabs["02_Animais"] if row["animal_id"] != "animal-42")
    assert created_animal["birth_date"] == "2024-01-02"
    assert created_animal["birth_date_precision"] == "DAY"
    assert created_store.tabs["06_Eventos"][-1]["event_type"] == "ANIMAL_CREATED"

    entered_store = MemoryStore()
    entered = make_processor(entered_store, "ANIMAL_ENTERED", {"display_identifier": "42"})
    preview = await receive(entered)
    await entered.confirm(42, preview.pending_id)
    assert entered_store.tabs["06_Eventos"][-1]["event_type"] == "ANIMAL_ENTERED"
    entry_projection = next(
        (row for row in entered_store.tabs["02_Animais"] if row.get("first_entry_date")),
        None,
    )
    assert entry_projection is not None, entered_store.tabs["02_Animais"]
    assert entry_projection["first_entry_date"] == entered_store.tabs["06_Eventos"][-1]["occurred_from"]
    assert entry_projection["entry_event_id"] == entered_store.tabs["06_Eventos"][-1]["event_id"]
    assert not entry_projection.get("birth_source")


@pytest.mark.asyncio
async def test_calving_confirmation_creates_outcome_calf_identifier_and_genealogy():
    store = MemoryStore()
    processor = make_processor(
        store,
        "CALVING",
        {
            "mother_identifier": "42",
            "outcomes": [
                {
                    "display_identifier": "391",
                    "live_status": "VIVO",
                    "sex": "FEMALE",
                    "outcome": "LIVE_CALF",
                    "birth_weight": 31.5,
                    "weight_unit": "kg",
                }
            ],
        },
    )

    preview = await receive(processor)
    assert "Parto" in preview.text and "391" in preview.text
    assert "31.5 kg" in preview.text
    assert "vínculo com a matriz" in preview.text
    result = await processor.confirm(42, preview.pending_id)

    assert result.text.startswith("✅")
    assert len(store.tabs["08A_Resultados_Parto"]) == 1
    assert store.tabs["08A_Resultados_Parto"][0]["birth_weight"] == 31.5
    assert store.tabs["02_Animais"][-1]["birth_source"] == store.tabs["06_Eventos"][-1]["event_id"]
    assert store.tabs["03_Identificadores_Animal"][-1]["value"] == "391"
    assert store.tabs["02A_Vinculos_Genealogicos"][0]["parent_animal_id"] == "animal-42"
    assert len(store.tabs["07_Evento_Animais"]) == 2


@pytest.mark.asyncio
async def test_calving_negative_liveness_status_does_not_create_calf_entity():
    store = MemoryStore()
    processor = make_processor(
        store,
        "CALVING",
        {
            "mother_identifier": "42",
            "outcomes": [
                {
                    "live_status": "NOT_LIVE",
                    "outcome": "STILLBORN",
                    "sex": "FEMALE",
                }
            ],
        },
    )

    preview = await receive(processor)
    result = await processor.confirm(42, preview.pending_id)

    assert result.text.startswith("✅")
    assert len(store.tabs["08A_Resultados_Parto"]) == 1
    assert len(store.tabs["02_Animais"]) == 1
    assert len(store.tabs["03_Identificadores_Animal"]) == 1
    assert not store.tabs["02A_Vinculos_Genealogicos"]


@pytest.mark.asyncio
async def test_calving_conflicting_liveness_values_require_clarification():
    store = MemoryStore()
    processor = make_processor(
        store,
        "CALVING",
        {
            "mother_identifier": "42",
            "outcomes": [
                {
                    "live_status": "NOT_LIVE",
                    "outcome": "LIVE_CALF",
                    "display_identifier": "391",
                }
            ],
        },
    )

    result = await receive(processor)

    assert "contradiz" in result.text
    assert not store.tabs["06_Eventos"]


@pytest.mark.asyncio
async def test_lot_transfer_closes_old_allocation_and_adds_new_one():
    store = MemoryStore()
    store.tabs["04_Grupos"].extend(
        [
            {"group_id": "g-old", "name": "3", "property_id": "farm", "valid_to": ""},
            {"group_id": "g-new", "name": "7", "property_id": "farm", "valid_to": ""},
        ]
    )
    store.tabs["05_Alocacoes"].append(
        {
            "allocation_id": "allocation-old",
            "property_id": "farm",
            "animal_id": "animal-42",
            "group_id": "g-old",
            "effective_from": "2026-01-01",
            "effective_to": "",
            "event_id": "old-event",
        }
    )
    processor = make_processor(store, "LOT_TRANSFER", {"animal_identifier": "42", "to_group_identifier": "7"})

    preview = await receive(processor)
    assert "Lote de origem: g-old" in preview.text
    result = await processor.confirm(42, preview.pending_id)

    assert result.text.startswith("✅")
    assert len(store.tabs["05_Alocacoes"]) == 3
    closed = store.tabs["05_Alocacoes"][1]
    new = store.tabs["05_Alocacoes"][2]
    assert closed["allocation_id"] == "allocation-old" and closed["effective_to"]
    assert new["group_id"] == "g-new" and not new["effective_to"]


@pytest.mark.asyncio
async def test_correction_is_append_only_and_history_hides_superseded_projection():
    store = MemoryStore()
    store.tabs["06_Eventos"].append(
        {
            "event_id": "prior-weight",
            "event_type": "WEIGHING",
            "property_id": "farm",
            "occurred_from": "2026-10-12",
            "command_id": "old-command",
            "processing_status": "APPLIED",
            "event_payload_json": '{"weight_kg": 512}',
            "correction_of": "",
        }
    )
    store.tabs["07_Evento_Animais"].append(
        {
            "event_id": "prior-weight",
            "animal_id": "animal-42",
            "property_id": "farm",
        }
    )
    processor = make_processor(
        store, "WEIGHING", {"animal_identifier": "42", "weight_kg": 510}, correction_of="prior-weight"
    )

    preview = await receive(processor)
    assert preview.pending_id
    result = await processor.confirm(42, preview.pending_id)
    history = await processor.receive(
        42,
        "/historico 42",
        source="text",
        message_id="chat:history",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )

    assert result.text.startswith("✅")
    assert len([row for row in store.tabs["06_Eventos"] if row["event_type"] == "WEIGHING"]) == 3
    assert store.tabs["06_Eventos"][-1]["correction_of"] == "prior-weight"
    assert "510 kg" in history.text
    assert "512 kg" not in history.text
    assert "corrige prior-weight" in history.text
    assert any(
        row.get("command_id") == "LOCAL:42:chat:history" and row.get("action") == "CONVERSATION_STATE"
        for row in store.tabs["99_Audit_Log"]
    )


@pytest.mark.asyncio
async def test_natural_weight_correction_links_latest_user_event_and_confirms_previous_value():
    store = MemoryStore()
    original = make_processor(store, "WEIGHING", {"animal_identifier": "42", "weight_kg": 87})
    first = await receive(original, "chat:weight-original")
    assert first.text.startswith("✅")
    original_event_id = store.tabs["06_Eventos"][-1]["event_id"]

    correction = make_processor(
        store,
        "WEIGHING",
        {"weight_kg": 78},
        correction_requested=True,
    )
    preview = await receive(correction, "chat:weight-correction")

    assert preview.pending_id == "LOCAL:42:chat:weight-correction"
    assert "Corrige o evento" in preview.text
    assert "Peso anterior: 87" in preview.text
    assert "78 kg" in preview.text
    corrected = await correction.confirm(42, preview.pending_id)

    assert corrected.text.startswith("✅")
    assert store.tabs["06_Eventos"][-1]["correction_of"] == original_event_id
    assert len([row for row in store.tabs["06_Eventos"] if row.get("event_id") == original_event_id]) == 2


@pytest.mark.asyncio
async def test_natural_correction_without_prior_event_asks_instead_of_creating_new_fact():
    store = MemoryStore()
    processor = make_processor(
        store,
        "WEIGHING",
        {"weight_kg": 78},
        correction_requested=True,
    )

    result = await receive(processor)

    assert "registro anterior compatível" in result.text
    assert not store.tabs["06_Eventos"]


@pytest.mark.asyncio
async def test_correction_for_unsupported_event_is_terminal_manual_review():
    store = MemoryStore()
    processor = make_processor(
        store,
        "CALVING",
        {"mother_identifier": "42", "outcomes": [{"outcome": "LIVE_CALF", "live_status": "VIVO"}]},
        correction_requested=True,
    )

    result = await receive(processor)

    assert "revisão manual" in result.text
    assert store.tabs["18_Conversas"][-1]["state"] == "MANUAL_REVIEW"
    assert await processor._latest_open_conversation_any_state(42) is None
    assert not store.tabs["06_Eventos"]


@pytest.mark.asyncio
async def test_correction_cannot_target_event_from_other_property_or_type():
    store = MemoryStore()
    store.tabs["06_Eventos"].append(
        {
            "event_id": "foreign",
            "event_type": "VACCINATION",
            "property_id": "other",
        }
    )
    processor = make_processor(
        store, "WEIGHING", {"animal_identifier": "42", "weight_kg": 510}, correction_of="foreign"
    )

    response = await receive(processor)

    assert "compatível para corrigir" in response.text
    assert len(store.tabs["06_Eventos"]) == 1
    assert store.tabs["06_Eventos"][0]["event_id"] == "foreign"


@pytest.mark.asyncio
async def test_group_vaccination_correction_cannot_change_target_group():
    store = MemoryStore()
    store.tabs["04_Grupos"].extend(
        [
            {"group_id": "group-7", "name": "7", "property_id": "farm", "valid_to": ""},
            {"group_id": "group-8", "name": "8", "property_id": "farm", "valid_to": ""},
        ]
    )
    store.tabs["06_Eventos"].append(
        {
            "event_id": "prior-group-vaccination",
            "event_type": "VACCINATION",
            "property_id": "farm",
            "processing_status": "APPLIED",
            "command_id": "old-command",
            "event_payload_json": '{"group_id":"group-7","group_identifier":"7","vaccine":"aftosa"}',
        }
    )
    store.tabs["07_Evento_Animais"].append(
        {
            "event_id": "prior-group-vaccination",
            "property_id": "farm",
            "group_id": "group-7",
            "target_basis": "GROUP_MEMBERSHIP",
        }
    )
    processor = make_processor(
        store,
        "VACCINATION",
        {"group_identifier": "8", "vaccine": "aftosa"},
        correction_of="prior-group-vaccination",
    )

    result = await receive(processor)

    assert "revisão manual" in result.text
    assert len(store.tabs["06_Eventos"]) == 1


@pytest.mark.asyncio
async def test_persisted_confirmation_binding_rejects_other_chat_and_message():
    store = MemoryStore()
    processor = make_processor(store, "VACCINATION", {"animal_identifier": "42", "vaccine": "vacina X"})
    preview = await receive(processor)
    await processor.bind_confirmation(42, preview.pending_id, chat_id=100, message_id=20)

    wrong_chat = await processor.confirm(42, preview.pending_id, source_chat_id=101, source_message_id=20)
    wrong_message = await processor.confirm(42, preview.pending_id, source_chat_id=100, source_message_id=21)
    accepted = await processor.confirm(42, preview.pending_id, source_chat_id=100, source_message_id=20)

    assert "não pertence a esta conversa" in wrong_chat.text
    assert "não pertence a esta mensagem" in wrong_message.text
    assert accepted.text.startswith("✅")
    assert len(store.tabs["06_Eventos"]) == 2


@pytest.mark.asyncio
async def test_missing_weight_is_clarified_then_applied_as_one_command():
    store = MemoryStore()
    clarification = InterpretationProposal(
        events=[
            EventProposal(
                event_type="WEIGHING", payload={"animal_identifier": "42"}, time_expression="hoje", confidence=0.98
            )
        ]
    )
    completed = InterpretationProposal(
        events=[
            EventProposal(
                event_type="WEIGHING",
                payload={"animal_identifier": "42", "weight_kg": 512},
                time_expression="hoje",
                confidence=0.98,
            )
        ]
    )

    class TwoStepInterpreter:
        def __init__(self):
            self.proposals = [clarification, completed]

        def interpret(self, _text):
            return self.proposals.pop(0)

    processor = MVPProcessor(
        local_settings(),
        store,
        TwoStepInterpreter(),
    )
    first = await receive(processor, "chat:1")
    second = await processor.receive(
        42,
        "512 quilos",
        source="text",
        message_id="chat:2",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )

    assert "Preciso esclarecer" in first.text
    assert second.text.startswith("✅")
    assert len(store.tabs["06_Eventos"]) == 2
    assert store.tabs["06_Eventos"][-1]["command_id"] == "LOCAL:42:chat:1"
    assert [row["state"] for row in store.tabs["18_Conversas"][-2:]] == ["NEEDS_CLARIFICATION", "COMPLETED"]


async def test_bare_hex_identifier_completes_an_authorized_clarification():
    store = MemoryStore()
    identifier = "0123456789ABCDEF0123"
    store.tabs["03_Identificadores_Animal"][0]["value"] = identifier
    clarification = InterpretationProposal(events=[EventProposal(
        event_type="HEALTH_OBSERVATION", payload={"observation": "manca"},
        time_expression="hoje", confidence=0.98,
    )])
    completed = InterpretationProposal(events=[EventProposal(
        event_type="HEALTH_OBSERVATION", payload={"animal_identifier": identifier, "observation": "manca"},
        time_expression="hoje", confidence=0.98,
    )])
    processor = MVPProcessor(
        local_settings(),
        store, SequenceInterpreter(clarification, completed),
    )

    first = await receive(processor)
    second = await processor.receive(
        42, identifier, source="text", message_id="chat:2",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None, source_reference=None,
    )

    assert "Preciso esclarecer" in first.text
    assert second.text.startswith("✅")
    assert store.tabs["06_Eventos"][-1]["command_id"] == "LOCAL:42:chat:1"
    assert store.tabs["18_Conversas"][-1]["state"] == "COMPLETED"




@pytest.mark.asyncio
async def test_multiple_clarifications_preserve_original_request_and_prior_answers():
    store = MemoryStore()
    proposals = [
        InterpretationProposal(
            events=[
                EventProposal(
                    event_type="WEIGHING",
                    payload={"animal_identifier": "42"},
                    time_expression="hoje",
                    confidence=0.98,
                )
            ]
        ),
        InterpretationProposal(
            events=[
                EventProposal(
                    event_type="WEIGHING",
                    payload={"animal_identifier": "42", "weight_kg": 512},
                    time_expression=None,
                    confidence=0.98,
                )
            ]
        ),
        InterpretationProposal(
            events=[
                EventProposal(
                    event_type="WEIGHING",
                    payload={"animal_identifier": "42", "weight_kg": 512},
                    time_expression="hoje",
                    confidence=0.98,
                )
            ]
        ),
    ]

    class CapturingInterpreter:
        def __init__(self):
            self.inputs = []

        def interpret(self, text):
            self.inputs.append(text)
            return proposals.pop(0)

    interpreter = CapturingInterpreter()
    processor = MVPProcessor(local_settings(), store, interpreter)
    first = await processor.receive(
        42,
        "A 184 pesou",
        source="text",
        message_id="chat:multi-1",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )
    second = await processor.receive(
        42,
        "512 quilos",
        source="text",
        message_id="chat:multi-2",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )
    third = await processor.receive(
        42,
        "hoje",
        source="text",
        message_id="chat:multi-3",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference=None,
    )

    assert "weight_kg" in first.text
    assert "time_expression" in second.text
    assert third.text.startswith("✅")
    assert "A 184 pesou" in interpreter.inputs[0]
    assert "A 184 pesou" in interpreter.inputs[1]
    assert "512 quilos" in interpreter.inputs[1]
    assert "A 184 pesou" in interpreter.inputs[2]
    assert "512 quilos" in interpreter.inputs[2]
    assert "Resposta do produtor: hoje" in interpreter.inputs[2]


@pytest.mark.asyncio
async def test_inferred_event_date_uses_original_message_across_midnight_clarification():
    store = MemoryStore()
    proposals = [
        InterpretationProposal(events=[
            EventProposal(
                event_type="HEALTH_OBSERVATION", payload={"animal_identifier": "42"}, confidence=0.98
            )
        ]),
        InterpretationProposal(events=[
            EventProposal(
                event_type="HEALTH_OBSERVATION",
                payload={"animal_identifier": "42", "observation": "mancando"},
                confidence=0.98,
            )
        ]),
    ]

    class TwoStepInterpreter:
        def interpret(self, _text):
            return proposals.pop(0)

    processor = MVPProcessor(
        local_settings(), store, TwoStepInterpreter()
    )
    await processor.receive(
        42,
        "o animal está mancando",
        source="text",
        message_id="m:1",
        captured_at="2026-10-13T02:59:00+00:00",
        device_timezone=None,
        source_reference=None,
    )
    await processor.receive(
        42,
        "está mancando",
        source="text",
        message_id="m:2",
        captured_at="2026-10-13T03:01:00+00:00",
        device_timezone=None,
        source_reference=None,
    )

    event = store.tabs["06_Eventos"][-1]
    assert event["occurred_from"] == "2026-10-12"
    assert event["time_source"] == "INFERRED_FROM_CAPTURE_DATE"


@pytest.mark.asyncio
async def test_relative_date_from_original_transcript_keeps_original_capture_anchor():
    store = MemoryStore()
    proposals = [
        InterpretationProposal(events=[
            EventProposal(
                event_type="HEALTH_OBSERVATION",
                payload={"animal_identifier": "42"},
                time_expression="hoje",
                confidence=0.98,
            )
        ]),
        InterpretationProposal(events=[
            EventProposal(
                event_type="HEALTH_OBSERVATION",
                payload={"animal_identifier": "42", "observation": "mancando"},
                time_expression="hoje",
                confidence=0.98,
            )
        ]),
    ]

    class TwoStepInterpreter:
        def interpret(self, _text):
            return proposals.pop(0)

    processor = MVPProcessor(
        local_settings(), store, TwoStepInterpreter()
    )
    await processor.receive(
        42,
        "hoje o animal está mancando",
        source="text",
        message_id="explicit:1",
        captured_at="2026-10-13T02:59:00+00:00",
        device_timezone=None,
        source_reference=None,
    )
    await processor.receive(
        42,
        "sim, está mancando",
        source="text",
        message_id="explicit:2",
        captured_at="2026-10-13T03:01:00+00:00",
        device_timezone=None,
        source_reference=None,
    )

    event = store.tabs["06_Eventos"][-1]
    assert event["occurred_from"] == "2026-10-12"
    assert event["time_source"] == "EXPLICIT"


@pytest.mark.asyncio
async def test_transfer_without_unique_active_allocation_requests_clarification_before_event():
    store = MemoryStore()
    store.tabs["04_Grupos"].append({"group_id": "g-new", "name": "7", "property_id": "farm", "valid_to": ""})
    processor = make_processor(store, "LOT_TRANSFER", {"animal_identifier": "42", "to_group_identifier": "7"})

    response = await receive(processor)

    assert "lote atual" in response.text
    assert not store.tabs["06_Eventos"]


@pytest.mark.asyncio
async def test_correction_with_allocation_effects_goes_to_manual_review_without_writes():
    store = MemoryStore()
    store.tabs["04_Grupos"].extend(
        [
            {"group_id": "g-old", "name": "3", "property_id": "farm", "valid_to": ""},
            {"group_id": "g-new", "name": "7", "property_id": "farm", "valid_to": ""},
        ]
    )
    store.tabs["05_Alocacoes"].append(
        {
            "allocation_id": "allocation-current",
            "property_id": "farm",
            "animal_id": "animal-42",
            "group_id": "g-old",
            "effective_from": "2026-01-01",
            "effective_to": "",
            "event_id": "prior-transfer",
        }
    )
    store.tabs["06_Eventos"].append(
        {
            "event_id": "prior-transfer",
            "event_type": "LOT_TRANSFER",
            "property_id": "farm",
            "processing_status": "APPLIED",
            "command_id": "old-command",
        }
    )
    store.tabs["07_Evento_Animais"].append(
        {
            "event_id": "prior-transfer",
            "animal_id": "animal-42",
            "property_id": "farm",
        }
    )
    processor = make_processor(
        store,
        "LOT_TRANSFER",
        {"animal_identifier": "42", "to_group_identifier": "7"},
        correction_of="prior-transfer",
    )

    response = await receive(processor)

    assert "revisão manual" in response.text
    assert store.tabs["18_Conversas"][-1]["state"] == "MANUAL_REVIEW"
    assert len(store.tabs["06_Eventos"]) == 1
    assert len(store.tabs["05_Alocacoes"]) == 1


@pytest.mark.asyncio
async def test_simple_weight_success_reply_contains_animal_weight_and_date():
    processor = make_processor(MemoryStore(), "WEIGHING", {"animal_identifier": "42", "weight_kg": 512})

    response = await receive(processor)

    assert response.text.startswith("✅")
    assert "Peso registrado" in response.text
    assert "animal 42" in response.text
    assert "512 kg" in response.text
    assert "12/10/2026" in response.text




async def request_export(processor, text="/exportar", user_id=42, source="text"):
    return await processor.receive(
        user_id,
        text,
        source=source,
        message_id="chat:export",
        captured_at=datetime(2026, 10, 13, tzinfo=UTC).isoformat(),
        device_timezone=None,
        source_reference="voice-export" if source == "voice" else None,
    )


@pytest.mark.parametrize("first_request", ["event", "export"])
async def test_export_and_event_writes_share_the_same_lock(monkeypatch, first_request):
    store = MemoryStore()
    processor = make_processor(store)
    blocked = asyncio.Event()
    release = asyncio.Event()
    competing_started = asyncio.Event()
    operations = []
    original_store = processor._store
    original_to_thread = asyncio.to_thread

    async def gated_store(method, *args, **kwargs):
        operations.append("event")
        if first_request == "event" and method == "append_target":
            # The initial event exists, but its target and applied status do not.
            blocked.set()
            await release.wait()
        return await original_store(method, *args, **kwargs)

    async def gated_to_thread(function, *args, **kwargs):
        if function is build_property_export:
            operations.append("export")
            if first_request == "export":
                blocked.set()
                await release.wait()
        return await original_to_thread(function, *args, **kwargs)

    monkeypatch.setattr(processor, "_store", gated_store)
    monkeypatch.setattr(asyncio, "to_thread", gated_to_thread)

    async def run_request(kind, *, competing=False):
        if competing:
            competing_started.set()
        if kind == "event":
            return await receive(processor)
        return await request_export(processor)

    first = asyncio.create_task(run_request(first_request))
    competing = None
    try:
        await asyncio.wait_for(blocked.wait(), timeout=10)
        before = list(operations)
        second_request = "export" if first_request == "event" else "event"
        competing = asyncio.create_task(run_request(second_request, competing=True))
        # The competing request runs to its first suspension before this wakes.
        await asyncio.wait_for(competing_started.wait(), timeout=10)
        assert operations == before
        assert not competing.done()
    finally:
        release.set()
        tasks = [first] if competing is None else [first, competing]
        results = await asyncio.wait_for(asyncio.gather(*tasks), timeout=10)

    exported = results[0] if first_request == "export" else results[1]
    recorded = results[1] if first_request == "export" else results[0]
    assert recorded.text.startswith("✅")
    assert store.tabs["06_Eventos"][-1]["processing_status"] == "APPLIED"
    assert len(store.tabs["07_Evento_Animais"]) == 1
    assert exported.document is not None
    workbook = load_workbook(BytesIO(exported.document.content))
    observations = [value for sheet in workbook for row in sheet.values for value in row]
    assert ("manca" in observations) == (first_request == "event")


@pytest.mark.parametrize("text", ["/exportar", "exportar planilha", "me manda a planilha"])
@pytest.mark.parametrize("source", ["text", "voice"])
async def test_authorized_export_returns_workbook_without_interpretation_or_persistent_writes(text, source):
    store = MemoryStore()
    store.tabs["02_Animais"][0]["species"] = "trusted-cattle"
    before = deepcopy(store.tabs)
    interpreter = SimpleNamespace(interpret=Mock(side_effect=AssertionError("export must bypass interpretation")))
    processor = MVPProcessor(local_settings(), store, interpreter)

    result = await request_export(processor, text, source=source)

    assert result.document is not None
    assert result.document.filename.endswith(".xlsx")
    assert result.pending_id is None
    workbook = load_workbook(BytesIO(result.document.content))
    assert workbook.sheetnames
    assert any("trusted-cattle" in row for sheet in workbook for row in sheet.iter_rows(values_only=True))
    assert store.tabs == before
    interpreter.interpret.assert_not_called()






@pytest.mark.asyncio
async def test_export_waits_for_in_progress_event_write_lock():
    store = MemoryStore()
    processor = MVPProcessor(
        local_settings(),
        store,
        SimpleNamespace(interpret=Mock()),
    )

    async with processor._writer_lock:
        export_task = asyncio.create_task(request_export(processor))
        await asyncio.sleep(0.02)
        assert not export_task.done()

    result = await export_task
    assert result.document is not None




async def test_export_unauthorized_user_cannot_read_or_interpret_data():
    store = SimpleNamespace(read_rows=Mock(side_effect=AssertionError("unauthorized read")))
    interpreter = SimpleNamespace(interpret=Mock(side_effect=AssertionError("unauthorized interpretation")))
    processor = MVPProcessor(local_settings(), store, interpreter)

    result = await request_export(processor, user_id=99)

    assert result.document is None
    assert result.text == "Acesso não autorizado."
    store.read_rows.assert_not_called()
    interpreter.interpret.assert_not_called()




async def test_export_uses_trusted_property_and_excludes_foreign_and_missing_property_rows():
    store = MemoryStore()
    store.tabs["02_Animais"][0]["species"] = "trusted-cattle"
    store.tabs["02_Animais"].extend([
        {"animal_id": "private-other-animal", "property_id": "other-farm", "species": "private-species"},
        {"animal_id": "unscoped-animal", "species": "unscoped-species"},
    ])
    store.tabs["06_Eventos"].extend([
        {
            "event_id": "trusted-event", "property_id": "farm",
            "event_type": "HEALTH_OBSERVATION", "processing_status": "APPLIED",
            "event_payload_json": '{"observation": "trusted-observation"}',
        },
        {
            "event_id": "private-other-event", "property_id": "other-farm",
            "event_type": "HEALTH_OBSERVATION", "processing_status": "APPLIED",
            "event_payload_json": '{"observation": "private-observation"}',
        },
    ])
    before = deepcopy(store.tabs)
    processor = make_processor(store)
    processor.interpreter = SimpleNamespace(interpret=Mock(side_effect=AssertionError("must not interpret")))

    result = await request_export(processor, "/exportar other-farm")

    assert result.document is not None
    workbook = load_workbook(BytesIO(result.document.content))
    cells = [value for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
    assert "trusted-cattle" in cells
    assert "trusted-observation" in cells
    assert "private-species" not in cells
    assert "private-observation" not in cells
    assert "unscoped-species" not in cells
    assert store.tabs == before


async def test_empty_property_export_is_valid_workbook_with_headers_and_no_writes():
    store = MemoryStore()
    store.tabs = {name: [] for name in TAB_HEADERS}
    before = deepcopy(store.tabs)
    processor = make_processor(store)

    result = await request_export(processor)

    assert result.document is not None
    workbook = load_workbook(BytesIO(result.document.content))
    assert workbook.sheetnames
    assert workbook.sheetnames == ["Resumo"]
    assert list(workbook["Resumo"].values)[0] == ("Informação", "Valor")
    assert ("Animais", 0) in list(workbook["Resumo"].values)
    assert ("Eventos", 0) in list(workbook["Resumo"].values)
    assert any("Ainda não há dados" in str(value) for row in workbook["Resumo"].values for value in row)
    assert store.tabs == before


async def test_export_read_failure_returns_retry_reply_without_document_or_writes():
    store = MemoryStore()
    before = deepcopy(store.tabs)
    store.read_rows = Mock(side_effect=OSError("sensitive provider error"))
    processor = make_processor(store)

    result = await request_export(processor)

    assert result.document is None
    assert "Não foi possível gerar a planilha" in result.text
    assert "sensitive" not in result.text
    assert store.tabs == before
